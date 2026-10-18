from io import BytesIO

import pytest

from src.data.file_parser import parse_bytes


def test_parse_txt():
    assert parse_bytes("resume.txt", "熟悉Python和PyTorch".encode("utf-8")) == "熟悉Python和PyTorch"


def test_parse_docx():
    import docx
    buf = BytesIO()
    d = docx.Document()
    d.add_paragraph("熟悉Python")
    d.add_paragraph("了解分布式训练")
    d.save(buf)
    text = parse_bytes("resume.docx", buf.getvalue())
    assert "Python" in text
    assert "分布式训练" in text


def test_parse_unsupported():
    with pytest.raises(ValueError):
        parse_bytes("a.xyz", b"data")


def test_parse_pdf():
    pdf = _make_minimal_pdf("Hello PDF text layer")
    assert "Hello PDF" in parse_bytes("a.pdf", pdf)


def _make_minimal_pdf(text):
    content = f"BT /F1 24 Tf 72 700 Td ({text}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref_pos = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF".encode()
    return bytes(out)
