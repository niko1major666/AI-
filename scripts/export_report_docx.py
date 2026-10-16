"""把 PROJECT_REPORT.md 导出为 Word（.docx）。"""
import re
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")

from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
MD = ROOT / "PROJECT_REPORT.md"
OUT = ROOT / "PROJECT_REPORT.docx"

CN_FONT = "微软雅黑"


def _set_font(run, size=None, bold=False, mono=False, color=None):
    run.bold = bold
    run.font.name = "Consolas" if mono else CN_FONT
    try:
        run._element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)
    except AttributeError:
        pass
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)


def _add_inline(par, text):
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for p in parts:
        if not p:
            continue
        if p.startswith("**") and p.endswith("**"):
            _set_font(par.add_run(p[2:-2]), bold=True)
        elif p.startswith("`") and p.endswith("`"):
            _set_font(par.add_run(p[1:-1]), mono=True, size=10)
        else:
            _set_font(par.add_run(p))


def _heading(doc, text, level):
    h = doc.add_heading(level=level)
    _set_font(h.add_run(text), size={1: 18, 2: 15, 3: 13}[level], bold=True,
              color=(0x1F, 0x3A, 0x5F))


def main():
    doc = Document()
    lines = MD.read_text(encoding="utf-8").splitlines()
    i = 0
    n = len(lines)
    while i < n:
        s = lines[i].strip()
        if not s:
            i += 1
            continue
        if s.startswith("```"):
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                p = doc.add_paragraph()
                _set_font(p.add_run(lines[i]), mono=True, size=9)
                i += 1
            i += 1
            continue
        if s.startswith("### "):
            _heading(doc, s[4:], 3)
        elif s.startswith("## "):
            _heading(doc, s[3:], 2)
        elif s.startswith("# "):
            _heading(doc, s[2:], 1)
        elif s.startswith("|"):
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                rows.append(lines[i].strip())
                i += 1
            data = []
            for r in rows:
                cells = [c.strip() for c in r.strip("|").split("|")]
                if all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    continue
                data.append(cells)
            if data:
                ncol = max(len(r) for r in data)
                table = doc.add_table(rows=len(data), cols=ncol)
                table.style = "Table Grid"
                for ri, row in enumerate(data):
                    for ci in range(ncol):
                        cell = table.cell(ri, ci)
                        cell.text = row[ci] if ci < len(row) else ""
                        for par in cell.paragraphs:
                            for run in par.runs:
                                run.font.name = CN_FONT
                                try:
                                    run._element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)
                                except AttributeError:
                                    pass
            continue
        elif s.startswith("> "):
            p = doc.add_paragraph()
            _add_inline(p, s[2:])
            for r in p.runs:
                r.font.italic = True
                r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        elif re.match(r"^\d+\.\s", s):
            p = doc.add_paragraph(style="List Number")
            _add_inline(p, re.sub(r"^\d+\.\s", "", s))
        elif s.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            _add_inline(p, s[2:])
        else:
            p = doc.add_paragraph()
            _add_inline(p, s)
        i += 1

    doc.save(OUT)
    print(f"已导出 → {OUT}")


if __name__ == "__main__":
    main()
