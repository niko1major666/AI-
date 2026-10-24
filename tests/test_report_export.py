from io import BytesIO
from pathlib import Path

from docx import Document

from src.app.report_export import report_docx, report_markdown, safe_filename
from src.matcher.gap_analyzer import SkillGapAnalyzer
from src.matcher.learning_path import build_learning_path
from src.skillkb.loader import Taxonomy

ROOT = Path(__file__).resolve().parents[1]
TAX = ROOT / "data" / "kb" / "skill_taxonomy.yaml"


def _report():
    analyzer = SkillGapAnalyzer(Taxonomy(str(TAX)))
    jd = [
        {"skill_id": "PROG.PYTHON", "name": "Python", "is_must": True, "tf": 2},
        {"skill_id": "FRAMEWORK.PYTORCH", "name": "PyTorch", "is_must": True},
        {"skill_id": "AITECH.RAG", "name": "RAG", "is_must": False},
    ]
    return analyzer.analyze(["Python"], jd, demand={"AITECH.RAG": 0.4})


def test_safe_filename_strips_illegal_chars():
    name = safe_filename("腾讯/云与智慧", '大模型"算法"工程师', "docx")
    assert name.startswith("腾讯_云与智慧_")
    assert name.endswith(".docx")
    for ch in '\\/:*?"<>|':
        assert ch not in name


def test_safe_filename_fallback():
    assert safe_filename("", "", "md") == "gap_report.md"
    assert len(safe_filename("x" * 200, "y" * 200, "md")) <= 64


def test_markdown_contains_report_structure():
    md = report_markdown(_report(), company="腾讯", job_title="AI 算法工程师",
                         user_skills=["Python"], generated_at="2026-01-01 00:00")
    assert "# 技能差距报告" in md
    assert "腾讯 · AI 算法工程师" in md
    assert "匹配度" in md and "2026-01-01 00:00" in md
    assert "P0 · 先学" in md and "PyTorch" in md
    assert "缺失技能清单" in md
    assert "| 技能 | 优先级 |" in md


def test_markdown_without_missing_says_all_covered():
    analyzer = SkillGapAnalyzer(Taxonomy(str(TAX)))
    rep = analyzer.analyze(["Python"], [{"skill_id": "PROG.PYTHON", "name": "Python",
                                         "is_must": True}])
    md = report_markdown(rep, company="X", job_title="Y")
    assert "已全部覆盖" in md


def test_markdown_attaches_learning_path():
    rep = _report()
    path = build_learning_path(rep.missing)
    md = report_markdown(rep, learning_path=path)
    assert "## 学习路径（规则版）" in md
    assert "阶段 1" in md


def test_docx_is_valid_document():
    rep = _report()
    path = build_learning_path(rep.missing)
    data = report_docx(rep, company="腾讯", job_title="AI 算法工程师",
                       user_skills=["Python"], learning_path=path)
    assert data[:2] == b"PK"                       # docx = zip
    doc = Document(BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "技能差距报告" in text
    assert "腾讯 · AI 算法工程师" in text
    assert "缺失技能清单" in text
    assert "学习路径（规则版）" in text
    # 清单表格：1 行表头 + 缺失项
    assert len(doc.tables) == 1
    assert len(doc.tables[0].rows) == len(rep.missing) + 1
