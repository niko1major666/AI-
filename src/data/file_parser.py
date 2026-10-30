"""文件解析：把上传的简历（PDF / DOCX / TXT / MD）解析为纯文本。

PDF 优先提取文字层；若为扫描件（无文字层），回退到 OCR（pymupdf 渲染 + rapidocr）。
"""
from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Optional, Union

from .privacy import UploadValidationError

_SUPPORTED = {".pdf", ".docx", ".txt", ".md", ".text"}


def parse_bytes(filename: str, data: bytes, *, max_pages: Optional[int] = None) -> str:
    """解析上传文件字节流为文本（用于 Streamlit file_uploader）。

    max_pages 仅对 PDF 生效：在提取文字/OCR 前拒绝超页文件，避免扫描件把
    整个请求拖入长时间 OCR。默认 None 表示不限制（脚本解析本地文件时可用）。
    """
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        text = _pdf_text(data, max_pages=max_pages)
        if len(text.strip()) < 10:
            text = _pdf_ocr(data, max_pages=max_pages)
        return text
    if suffix == ".docx":
        import docx
        return _docx_to_text(docx.Document(BytesIO(data)))
    if suffix in (".txt", ".md", ".text", ""):
        return data.decode("utf-8", errors="ignore").strip()
    raise ValueError(f"不支持的文件类型：{suffix}（支持 {', '.join(sorted(_SUPPORTED))}）")


def parse_file(path: Union[str, Path]) -> str:
    """解析本地文件路径为文本。"""
    p = Path(path)
    return parse_bytes(p.name, p.read_bytes())


def _pdf_text(data: bytes, *, max_pages: Optional[int] = None) -> str:
    from pypdf import PdfReader
    reader = PdfReader(BytesIO(data))
    _check_page_limit(len(reader.pages), max_pages)
    return "\n".join((page.extract_text() or "") for page in reader.pages).strip()


def _pdf_ocr(data: bytes, *, max_pages: Optional[int] = None) -> str:
    """扫描件兜底：pymupdf 渲染页面 + rapidocr OCR。"""
    try:
        import fitz  # pymupdf
        import numpy as np
        from rapidocr_onnxruntime import RapidOCR
    except Exception as e:  # noqa: BLE001
        return f"[OCR 不可用：{e}。请确认 PDF 含文字层，或安装 pymupdf/rapidocr_onnxruntime]"

    try:
        engine = RapidOCR()
        doc = fitz.open(stream=data, filetype="pdf")
        _check_page_limit(len(doc), max_pages)
        lines = []
        for page in doc:
            pix = page.get_pixmap(dpi=200)
            img = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)
            img = img[:, :, :3][:, :, ::-1]  # RGB -> BGR
            result, _ = engine(img)
            if result:
                lines.append("\n".join(item[1] for item in result))
        return "\n".join(lines).strip()
    except Exception as e:  # noqa: BLE001
        return f"[OCR 失败：{e}]"


def _check_page_limit(pages: int, max_pages: Optional[int]) -> None:
    if max_pages and pages > max_pages:
        raise UploadValidationError(f"PDF 页数过多：{pages} 页（上限 {max_pages} 页）")


def _docx_to_text(doc) -> str:
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" ".join(cell.text for cell in row.cells))
    return "\n".join(parts).strip()

